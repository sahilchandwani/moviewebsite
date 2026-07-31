#!/usr/bin/env python3
"""One-off: convert a salons CSV into a nicely formatted .xlsx so it can be
handed to another Claude conversation (or a person) to draft messages from.
Not part of the normal 3-stage pipeline.

Usage: python build_excel.py [src.csv] [out.xlsx]
Defaults to data/salons.csv -> data/salons.xlsx.
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "data/salons.csv"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data/salons.xlsx"

HEADER_LABELS = {
    "handle": "Instagram Handle",
    "name": "Salon / Display Name",
    "bio": "Bio",
    "website": "Website / Link",
    "phone": "Phone",
    "whatsapp": "WhatsApp",
    "email": "Email",
}

with open(SRC, newline="", encoding="utf-8") as f:
    rows = list(csv.DictReader(f))

fieldnames = list(rows[0].keys()) if rows else list(HEADER_LABELS.keys())

wb = Workbook()
ws = wb.active
ws.title = "Salons"

header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
body_font = Font(name="Arial")
wrap = Alignment(vertical="top", wrap_text=True)

for col_idx, field in enumerate(fieldnames, start=1):
    cell = ws.cell(row=1, column=col_idx, value=HEADER_LABELS.get(field, field.title()))
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(vertical="center")

for row_idx, row in enumerate(rows, start=2):
    for col_idx, field in enumerate(fieldnames, start=1):
        value = row.get(field, "")
        # Prefix @ on the handle for readability.
        if field == "handle" and value:
            value = f"@{value}"
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = body_font
        cell.alignment = wrap

# Reasonable column widths per field, not just uniform.
widths = {
    "handle": 26, "name": 28, "bio": 55, "website": 30,
    "phone": 18, "whatsapp": 18, "email": 26,
}
for col_idx, field in enumerate(fieldnames, start=1):
    ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(field, 20)

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{len(rows) + 1}"

wb.save(OUT)
print(f"Wrote {len(rows)} row(s) to {OUT}")
